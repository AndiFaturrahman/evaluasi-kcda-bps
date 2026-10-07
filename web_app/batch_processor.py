import os
import glob
import zipfile
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from web_app.evaluator import extract_pdf_metadata, analyze_defects
from web_app.pdf_generator import generate_pdf_report

def populate_district_sheet(ws, meta, defects):
    region = meta.get("region", "Wilayah")
    year = meta.get("year", 2026)
    title = meta.get("title", f"Kecamatan {region} Dalam Angka {year}")

    font_header = Font(name='Calibri', size=11, bold=True)
    font_bold = Font(name='Calibri', size=11, bold=True)
    font_regular = Font(name='Calibri', size=11, bold=False)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )
    header_fill = PatternFill(start_color='E8EEF5', end_color='E8EEF5', fill_type='solid')

    headers = [
        ("No", 5), 
        ("Judul Publikasi", 35), 
        ("Periode Terbit", 14), 
        ("Bahasa", 18),
        ("Tanggal Rilis", 14), 
        ("Tanggal Periksa", 14),
        ("Rilis On Time Dan Softcopy Tersedia di web", 22), 
        ("Rilis Tidak On Time", 16),
        ("Realisasi Tanggal Rilis", 16), 
        ("Tidak Rilis", 14),
        ("Keterangan", 40), 
        ("Alasan", 90)
    ]

    ws.column_dimensions['A'].width = 3
    for idx, (h_title, w) in enumerate(headers, start=2):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = w
        
        c3 = ws.cell(row=3, column=idx, value=h_title)
        c3.font = font_header
        c3.alignment = align_center
        c3.border = thin_border
        c3.fill = header_fill
        
        c4 = ws.cell(row=4, column=idx, value=f"({idx-1})")
        c4.font = font_header
        c4.alignment = align_center
        c4.border = thin_border
        c4.fill = header_fill

    kabupaten_display = meta.get("kabupaten")
    if not kabupaten_display:
        try:
            from evaluator import detect_kabupaten_name
            kabupaten_display = detect_kabupaten_name(pdf_path=None, pages_text={}, region_name=region, pub_title=title, is_kabupaten=meta.get("is_kabupaten", False))
        except Exception:
            kabupaten_display = f"Kabupaten {region}"
    c2 = ws.cell(row=2, column=2, value=kabupaten_display)
    c2.font = Font(name='Calibri', size=12, bold=True)

    ws.row_dimensions[5].height = 45
    ws.cell(row=5, column=2, value=1).alignment = align_center
    ws.cell(row=5, column=3, value=title).alignment = align_left
    ws.cell(row=5, column=4, value="Tahunan").alignment = align_center
    ws.cell(row=5, column=5, value="Indonesia dan Inggris").alignment = align_center
    ws.cell(row=5, column=6, value=f"September {year}").alignment = align_center
    ws.cell(row=5, column=7, value=f"29 September {year}").alignment = align_center
    ws.cell(row=5, column=8, value="√").alignment = align_center
    ws.cell(row=5, column=9, value="").alignment = align_center
    ws.cell(row=5, column=10, value=f"September {year}").alignment = align_center
    ws.cell(row=5, column=11, value="").alignment = align_center

    for c in range(2, 14):
        cell = ws.cell(row=5, column=c)
        cell.font = font_regular
        cell.border = thin_border

    curr_row = 5
    first_section = True
    
    for section_header, items in defects.items():
        if not items:
            continue
            
        if not first_section:
            c_sec = ws.cell(row=curr_row, column=12, value=section_header)
            c_sec.font = font_bold
            c_sec.alignment = align_left
            c_sec.border = thin_border
            ws.cell(row=curr_row, column=13).border = thin_border
            ws.row_dimensions[curr_row].height = 25
            curr_row += 1
        else:
            c_sec = ws.cell(row=5, column=12, value=section_header)
            c_sec.font = font_bold
            c_sec.alignment = align_left
            c_sec.border = thin_border
            ws.cell(row=5, column=13).border = thin_border
            first_section = False
            curr_row = 6

        for item in items:
            c_bullet = ws.cell(row=curr_row, column=12, value="-")
            c_bullet.alignment = align_center
            c_bullet.font = font_regular
            c_bullet.border = thin_border
            
            c_desc = ws.cell(row=curr_row, column=13, value=item)
            c_desc.font = font_regular
            c_desc.alignment = align_left
            c_desc.border = thin_border
            
            line_estimate = max(1, len(item) // 75 + item.count("\n") + 1)
            ws.row_dimensions[curr_row].height = max(28, min(line_estimate * 18, 140))
            
            curr_row += 1


def generate_master_batch_excel(eval_results, output_excel_path, regency_title="Kolektif Publikasi Kecamatan"):
    wb = openpyxl.Workbook()
    # Sheet 1: Rekapitulasi Wilayah
    ws_rekap = wb.active
    ws_rekap.title = "Rekapitulasi Wilayah"

    font_title = Font(name='Calibri', size=13, bold=True, color='0F172A')
    font_sub = Font(name='Calibri', size=10, italic=True, color='475569')
    font_h = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    font_row = Font(name='Calibri', size=10, bold=False)
    font_row_bold = Font(name='Calibri', size=10, bold=True)
    
    fill_h = PatternFill(start_color='004080', end_color='004080', fill_type='solid')
    fill_fatal = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    fill_rev = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    fill_clean = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    ws_rekap.cell(row=2, column=2, value=f"REKAPITULASI HASIL AUDIT KEPATUHAN FORMAT PUBLIKASI KCDA {eval_results[0]['meta'].get('year', 2026)}").font = font_title
    ws_rekap.cell(row=3, column=2, value=f"Wilayah: {regency_title} | Berdasarkan Pedoman Pembuatan Publikasi BPS 2023 & Standar Evaluasi 2026").font = font_sub

    rekap_headers = [
        ("No", 6),
        ("Kecamatan", 20),
        ("Judul Publikasi", 35),
        ("Nomor Katalog", 18),
        ("Nomor ISSN", 16),
        ("Total Temuan", 14),
        ("Status Kelayakan", 22),
        ("Integritas Wilayah (Kover vs Isi)", 28),
        ("Rincian Temuan Kritis / Saran Estetika", 50)
    ]

    ws_rekap.row_dimensions[5].height = 28
    for c_idx, (h_name, w) in enumerate(rekap_headers, start=2):
        col_letter = openpyxl.utils.get_column_letter(c_idx)
        ws_rekap.column_dimensions[col_letter].width = w
        cell = ws_rekap.cell(row=5, column=c_idx, value=h_name)
        cell.font = font_h
        cell.fill = fill_h
        cell.alignment = align_center
        cell.border = thin_border

    for r_idx, res in enumerate(eval_results, start=6):
        meta = res["meta"]
        defects = res["defects"]
        tot_d = res["total_defects"]
        d_mismatch = meta.get("district_mismatch_info", {})
        is_mismatch = d_mismatch.get("is_mismatch", False)

        if is_mismatch:
            status_str = "REVISI TOTAL (FATAL)"
            row_fill = fill_fatal
            font_status = Font(name='Calibri', size=10, bold=True, color='B91C1C')
            clone_status_str = f"FATAL: Beda Wilayah (Kover {d_mismatch.get('cover_district')} vs Isi {d_mismatch.get('dominant_inner_district') or d_mismatch.get('catalog_district')})"
        elif tot_d == 0:
            status_str = "SESUAI STANDAR"
            row_fill = fill_clean
            font_status = Font(name='Calibri', size=10, bold=True, color='15803D')
            clone_status_str = "Sesuai (Kover & Isi Sinkron)"
        else:
            status_str = "PERLU REVISI"
            row_fill = fill_rev
            font_status = Font(name='Calibri', size=10, bold=True, color='B45309')
            clone_status_str = "Sesuai (Kover & Isi Sinkron)"

        # Critical summary
        crit_list = []
        if is_mismatch:
            crit_list.append("FATAL Beda Wilayah")
        if meta.get("dummy_camera_detected") or meta.get("dummy_figures"):
            crit_list.append("Placeholder Ikon Kamera")
        if meta.get("issn_cross_page_inconsistent"):
            crit_list.append("Inkonsistensi ISSN")
        if any("[SARAN ESTETIKA" in x for items in defects.values() for x in items):
            crit_list.append("Saran Estetika Foto/Tabel")
        crit_str = ", ".join(crit_list) if crit_list else "Penyelarasan Standar BPS"

        ws_rekap.row_dimensions[r_idx].height = 24
        ws_rekap.cell(row=r_idx, column=2, value=r_idx-5).alignment = align_center
        ws_rekap.cell(row=r_idx, column=3, value=meta.get("region", "-")).alignment = align_left
        ws_rekap.cell(row=r_idx, column=4, value=meta.get("title", "-")).alignment = align_left
        ws_rekap.cell(row=r_idx, column=5, value=meta.get("catalog", "-")).alignment = align_center
        ws_rekap.cell(row=r_idx, column=6, value=meta.get("issn", "-")).alignment = align_center
        ws_rekap.cell(row=r_idx, column=7, value=tot_d).alignment = align_center
        
        c_stat = ws_rekap.cell(row=r_idx, column=8, value=status_str)
        c_stat.alignment = align_center
        c_stat.font = font_status
        c_stat.fill = row_fill

        c_clone = ws_rekap.cell(row=r_idx, column=9, value=clone_status_str)
        c_clone.alignment = align_left
        if is_mismatch:
            c_clone.font = font_status
            c_clone.fill = row_fill

        ws_rekap.cell(row=r_idx, column=10, value=crit_str).alignment = align_left

        for c_idx in range(2, 11):
            cell = ws_rekap.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            if c_idx not in [8, 9]:
                cell.font = font_row

    # Sheets 2+: Individual District Sheets
    used_titles = {"Rekapitulasi Wilayah"}
    for res in eval_results:
        meta = res["meta"]
        defects = res["defects"]
        reg_title = f"{meta.get('region', 'Kec')} {meta.get('year', 2026)}"[:30]
        if reg_title in used_titles:
            reg_title = f"{reg_title[:25]}_{len(used_titles)}"
        used_titles.add(reg_title)

        ws_dist = wb.create_sheet(title=reg_title)
        populate_district_sheet(ws_dist, meta, defects)

    os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)
    wb.save(output_excel_path)
    return output_excel_path


def create_batch_zip(file_items, output_zip_path):
    os.makedirs(os.path.dirname(output_zip_path), exist_ok=True)
    with zipfile.ZipFile(output_zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        for f_path, arc_name in file_items:
            if os.path.exists(f_path):
                zf.write(f_path, arc_name)
    return output_zip_path
