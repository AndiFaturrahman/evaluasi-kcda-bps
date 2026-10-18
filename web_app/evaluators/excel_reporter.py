"""
Modul Pembuatan Lembar Kerja Evaluasi Excel (.xlsx) Resmi BPS (Phase 4)
Menghasilkan tabel evaluasi sesuai format baku pemeriksaan publikasi BPS,
lengkap dengan kode kolom, status rilis, identitas satker, dan basis master wilayah.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from .common import detect_kabupaten_name, WILAYAH_METADATA

def generate_excel_report(meta, defects, output_path, base_template_path=None):
    region = meta["region"]
    year = meta["year"]
    title = meta["title"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{region} {year}"

    font_header = Font(name='Calibri', size=11, bold=True)
    font_bold = Font(name='Calibri', size=11, bold=True)
    font_regular = Font(name='Calibri', size=11, bold=False)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )
    header_fill = PatternFill(start_color='E8EEF5', end_color='E8EEF5', fill_type='solid')

    headers = [
        ("No", 5), 
        ("Judul Publikasi", 35), 
        ("Periode Terbit", 14), 
        ("Bahasa", 18),
        ("Tanggal Rilis", 14), 
        ("Tanggal Periksa", 14),
        ("Rilis On Time Dan Softcopy Tersedia di web", 22), 
        ("Rilis Tidak On Time", 16),
        ("Realisasi Tanggal Rilis", 16), 
        ("Tidak Rilis", 14),
        ("Keterangan", 40), 
        ("Alasan", 90)
    ]

    ws.column_dimensions['A'].width = 3
    for idx, (h_title, w) in enumerate(headers, start=2):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = w
        
        c3 = ws.cell(row=3, column=idx, value=h_title)
        c3.font = font_header
        c3.alignment = align_center
        c3.border = thin_border
        c3.fill = header_fill
        
        c4 = ws.cell(row=4, column=idx, value=f"({idx-1})")
        c4.font = font_header
        c4.alignment = align_center
        c4.border = thin_border
        c4.fill = header_fill

    kabupaten_display = meta.get("kabupaten")
    if not kabupaten_display:
        kabupaten_display = detect_kabupaten_name(pdf_path=None, pages_text={}, region_name=region, pub_title=title, is_kabupaten=meta.get("is_kabupaten", False))
    c2 = ws.cell(row=2, column=2, value=kabupaten_display)
    c2.font = Font(name='Calibri', size=12, bold=True)


    # Cantumkan referensi master dataset wilayah BPS & Kemendagri
    v_src = WILAYAH_METADATA.get("source", "Kemendagri & BPS RI Master Wilayah 2024")
    v_ver = WILAYAH_METADATA.get("version", "v2024.1")
    v_tot = WILAYAH_METADATA.get("total_districts", 7288)
    ws.cell(row=1, column=2, value=f"Referensi Basis Data: {v_src} ({v_ver} | {v_tot} Kecamatan, 514 Kab/Kota)").font = Font(name='Calibri', size=9, italic=True, color='475569')

    ws.row_dimensions[5].height = 45
    ws.cell(row=5, column=2, value=1).alignment = align_center
    ws.cell(row=5, column=3, value=title).alignment = align_left
    ws.cell(row=5, column=4, value="Tahunan").alignment = align_center
    ws.cell(row=5, column=5, value="Indonesia dan Inggris").alignment = align_center
    ws.cell(row=5, column=6, value=f"September {year}").alignment = align_center
    ws.cell(row=5, column=7, value=f"29 September {year}").alignment = align_center
    raw_tpl_meta = meta.get("raw_template_info", {})
    if raw_tpl_meta.get("is_pure_template"):
        ws.cell(row=5, column=8, value="").alignment = align_center
        ws.cell(row=5, column=9, value="").alignment = align_center
        ws.cell(row=5, column=10, value="Tidak Rilis").alignment = align_center
        ws.cell(row=5, column=11, value="√").alignment = align_center
    else:
        ws.cell(row=5, column=8, value="√").alignment = align_center
        ws.cell(row=5, column=9, value="").alignment = align_center
        ws.cell(row=5, column=10, value=f"September {year}").alignment = align_center
        ws.cell(row=5, column=11, value="").alignment = align_center

    for c in range(2, 14):
        cell = ws.cell(row=5, column=c)
        cell.font = font_regular
        cell.border = thin_border

    curr_row = 5
    first_section = True
    
    for section_header, items in defects.items():
        if not items:
            continue
            
        if not first_section:
            c_sec = ws.cell(row=curr_row, column=12, value=section_header)
            c_sec.font = font_bold
            c_sec.alignment = align_left
            c_sec.border = thin_border
            ws.cell(row=curr_row, column=13).border = thin_border
            ws.row_dimensions[curr_row].height = 25
            curr_row += 1
        else:
            c_sec = ws.cell(row=5, column=12, value=section_header)
            c_sec.font = font_bold
            c_sec.alignment = align_left
            c_sec.border = thin_border
            ws.cell(row=5, column=13).border = thin_border
            first_section = False
            curr_row = 6

        for item in items:
            c_bullet = ws.cell(row=curr_row, column=12, value="-")
            c_bullet.alignment = align_center
            c_bullet.font = font_regular
            c_bullet.border = thin_border
            
            c_desc = ws.cell(row=curr_row, column=13, value=item)
            c_desc.font = font_regular
            c_desc.alignment = align_left
            c_desc.border = thin_border
            
            line_estimate = max(1, len(item) // 75 + item.count("\n") + 1)
            ws.row_dimensions[curr_row].height = max(28, min(line_estimate * 18, 140))
            
            curr_row += 1

    wb.save(output_path)
    return output_path
